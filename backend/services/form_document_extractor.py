"""Structured extraction for uploaded forms and consents."""

from __future__ import annotations

import io
from typing import Any


def _pdf_widget_type(field_type: str, flags: int) -> str:
    if field_type == "/Tx":
        return "text"

    if field_type == "/Ch":
        return "select"

    if field_type == "/Sig":
        return "signature"

    if field_type == "/Btn":
        if flags & (1 << 15):
            return "radio"

        if flags & (1 << 16):
            return "button"

        return "checkbox"

    return "text"


def extract_pdf_structure(data: bytes) -> dict[str, Any]:
    """Extract visible text and embedded AcroForm widgets."""
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(data))

    pages: list[dict[str, Any]] = []
    widgets: list[dict[str, Any]] = []

    for page_index, page in enumerate(reader.pages):
        page_number = page_index + 1

        pages.append({
            "page": page_number,
            "text": page.extract_text() or "",
        })

        annotations = page.get("/Annots") or []

        for annotation_ref in annotations:
            try:
                annotation = annotation_ref.get_object()
            except Exception:
                continue

            if annotation.get("/Subtype") != "/Widget":
                continue

            parent_ref = annotation.get("/Parent")

            try:
                parent = (
                    parent_ref.get_object()
                    if parent_ref is not None
                    else annotation
                )
            except Exception:
                parent = annotation

            field_name = (
                annotation.get("/T")
                or parent.get("/T")
                or ""
            )

            field_type = str(
                annotation.get("/FT")
                or parent.get("/FT")
                or ""
            )

            flags = int(
                annotation.get("/Ff")
                or parent.get("/Ff")
                or 0
            )

            rect = annotation.get("/Rect") or []

            bounding_box = None

            if len(rect) == 4:
                try:
                    bounding_box = [
                        float(value)
                        for value in rect
                    ]
                except Exception:
                    bounding_box = None

            raw_options = (
                annotation.get("/Opt")
                or parent.get("/Opt")
                or []
            )

            options = []

            for option in raw_options:
                if isinstance(option, (list, tuple)):
                    options.append(
                        " | ".join(str(part) for part in option)
                    )
                else:
                    options.append(str(option))

            widgets.append({
                "page": page_number,
                "field_name": str(field_name),
                "field_type": _pdf_widget_type(
                    field_type,
                    flags,
                ),
                "pdf_field_type": field_type,
                "required": bool(flags & 2),
                "options": options,
                "value": str(
                    annotation.get("/V")
                    or parent.get("/V")
                    or ""
                ),
                "bounding_box": bounding_box,
                "source": "pdf_widget",
            })

    return {
        "document_type": "pdf",
        "page_count": len(reader.pages),
        "pages": pages,
        "widgets": widgets,
        "widget_count": len(widgets),
        "has_fillable_fields": bool(widgets),
    }


def extract_docx_structure(data: bytes) -> dict[str, Any]:
    """Extract paragraphs and tables from DOCX files."""
    from docx import Document

    document = Document(io.BytesIO(data))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text and paragraph.text.strip()
    ]

    tables = []

    for table_index, table in enumerate(document.tables):
        rows = []

        for row in table.rows:
            rows.append([
                cell.text.strip()
                for cell in row.cells
            ])

        tables.append({
            "table_index": table_index,
            "rows": rows,
        })

    return {
        "document_type": "docx",
        "paragraphs": paragraphs,
        "tables": tables,
        "has_fillable_fields": False,
    }


def extract_form_structure(
    filename: str,
    data: bytes,
) -> dict[str, Any]:
    name = (filename or "").lower()

    if name.endswith(".pdf"):
        return extract_pdf_structure(data)

    if name.endswith(".docx"):
        return extract_docx_structure(data)

    if name.endswith(".txt"):
        return {
            "document_type": "txt",
            "text": data.decode(
                "utf-8",
                errors="replace",
            ),
            "has_fillable_fields": False,
        }

    raise ValueError(
        "Unsupported document type. Upload PDF, DOCX, or TXT."
    )
